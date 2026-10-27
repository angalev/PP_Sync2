"""Синтетические «боевой файл» и «копия с изменениями» для локальных прогонов.

Повторяет структуру реальной ПП: двухуровневая шапка в строках 6-7, коды
столбцов в строке 9, данные с 10-й, ключ в столбце U (код 39), трансгаз в A
(код 7), объект в F (код 14).
Общего доступа и макросов тут нет — они проверяются только на реальном файле.

    python -m tools.make_test_workbook            # оба файла в tests/data
"""

from __future__ import annotations

import datetime as _dt
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

SHEET = "ПП 2026"
FIRST_DATA_ROW = 10
LAST_COLUMN = 95  # до CQ

TRANSGAZES = [
    ("Волгоград ГТ", 20),
    ("Краснодар ГТ", 15),
    ("Санкт-Петербург ГТ", 10),
    ("Казань ГТ", 8),
    ("Югорск ГТ", 3),
]

# Двухуровневая шапка: (буква, родитель в строке 6, уточнение в строке 7)
HEADERS = {
    "A": ("Общие сведения", "Наименование трансгаза"),
    "F": ("Общие сведения", "Наименование объекта ремонта"),
    "U": ("Коды", "Сводный код объекта ремонта (X/Y/Z)"),
    "CJ": ("ОТД", "Дата передачи на согласование ТО"),
    "CK": ("ОТД", "Дата согласования ТО"),
    "CL": ("ОТД", "Номер техакта"),
    "CM": ("ОТД", "Примечание ОТД"),
}


# Коды столбцов в строке 9 — как в реальной ПП: 39 — ключ X/Y/Z, 7 — трансгаз,
# 14 — объект, 313 — расчётный столбец (есть в Config.formula_codes).
CODES = {
    "A": 7,
    "F": 14,
    "U": 39,
    "CJ": 271,
    "CK": 313,
    "CL": 316,
    "CM": 325,
    get_column_letter(LAST_COLUMN): 999,
}


def build(path: Path) -> Path:
    """Собрать «боевой» файл."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET

    sheet["A1"] = "Производственная программа 2026 (ТЕСТОВАЯ КОПИЯ)"

    for letter, (parent, child) in HEADERS.items():
        sheet[f"{letter}6"] = parent
        sheet[f"{letter}7"] = child

    # Строка 9 — коды столбцов: по ним программа находит столбцы
    for letter, code in CODES.items():
        sheet[f"{letter}9"] = code

    row = FIRST_DATA_ROW
    counter = 0
    for transgaz, count in TRANSGAZES:
        for number in range(1, count + 1):
            counter += 1
            key = f"97575/{counter % 17 + 1}/{100000 + counter * 37}"
            sheet[f"A{row}"] = transgaz
            sheet[f"F{row}"] = f"{transgaz.split()[0]} — объект {number}"
            sheet[f"U{row}"] = key
            # Часть строк уже заполнена — чтобы получить статус «перезапись»
            if counter % 5 == 0:
                sheet[f"CJ{row}"] = _dt.datetime(2026, 3, 1) + _dt.timedelta(  # noqa: DTZ001
                    days=counter
                )
                sheet[f"CJ{row}"].number_format = "DD.MM.YYYY"
                sheet[f"CM{row}"] = "старое примечание"
            # Одна ячейка с формулой — проверка статуса «формула в ячейке»
            if counter == 3:
                sheet[f"CK{row}"] = f"=CJ{row}+7"
            row += 1

    for letter in ("A", "F", "U"):
        sheet.column_dimensions[letter].width = 30
    for letter in ("CJ", "CK", "CL", "CM"):
        sheet.column_dimensions[letter].width = 18

    sheet.cell(row=6, column=LAST_COLUMN, value="Прочее")
    sheet.cell(row=7, column=LAST_COLUMN, value=get_column_letter(LAST_COLUMN))

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def build_changed(target: Path, source: Path) -> dict[str, str]:
    """Сделать из «боевого» файла копию с изменениями.

    Возвращает описание внесённых случаев: адрес -> что проверяем.
    """
    workbook = load_workbook(target)
    sheet = workbook[SHEET]

    cases: dict[str, str] = {}
    rows = [r for r in range(FIRST_DATA_ROW, sheet.max_row + 1) if sheet[f"U{r}"].value]

    # 1. в боевом пусто, в копии значение -> «к применению»
    first = rows[0]
    sheet[f"CJ{first}"] = _dt.datetime(2026, 6, 30)  # noqa: DTZ001
    sheet[f"CJ{first}"].number_format = "DD.MM.YYYY"
    cases[f"CJ{first}"] = "к применению (в боевом пусто)"

    # 2. в боевом значение, в копии другое -> «перезапись»
    filled = rows[4]
    sheet[f"CJ{filled}"] = _dt.datetime(2026, 7, 15)  # noqa: DTZ001
    sheet[f"CJ{filled}"].number_format = "DD.MM.YYYY"
    cases[f"CJ{filled}"] = "перезапись"

    # 3. текст в пустую ячейку
    text_row = rows[1]
    sheet[f"CL{text_row}"] = "ТА-2026/117"
    cases[f"CL{text_row}"] = "текст в пустую ячейку"

    # 4. в боевом формула -> должно заблокироваться
    formula_row = rows[2]
    sheet[f"CK{formula_row}"] = _dt.datetime(2026, 8, 1)  # noqa: DTZ001
    cases[f"CK{formula_row}"] = "поверх формулы (ожидаем блокировку)"

    # 5. в копии пусто, в боевом было -> «очистка»
    sheet[f"CM{filled}"] = None
    cases[f"CM{filled}"] = "очистка"

    # 6. строки нет в боевом файле
    new_row = sheet.max_row + 1
    sheet[f"A{new_row}"] = "Волгоград ГТ"
    sheet[f"F{new_row}"] = "Объект, которого нет в боевом"
    sheet[f"U{new_row}"] = "99999/99/999999"
    sheet[f"CL{new_row}"] = "ТА-НОВАЯ"
    cases[f"CL{new_row}"] = "строки нет в боевом"

    source.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(source)
    return cases


def main() -> None:
    base = Path(sys.argv[1]).parent if len(sys.argv) > 1 else Path("tests/data")
    target = base / "ПП_боевой.xlsx"
    source = base / "ПП_копия_с_изменениями.xlsx"
    build(target)
    cases = build_changed(target, source)
    print(f"Боевой файл : {target.resolve()}")
    print(f"Копия       : {source.resolve()}")
    print(f"Различий заложено: {len(cases)}")
    for address, case in cases.items():
        print(f"  {address:<8} {case}")


if __name__ == "__main__":
    main()
