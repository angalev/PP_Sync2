"""Проверка полного цикла записи в песочнице.

Собирает пару файлов (боевой + копия с изменениями), прогоняет SyncSession,
затем независимо перечитывает результат через openpyxl и сверяет с ожидаемым.

    python -m tools.apply_check
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from ppsync.applier import SyncSession
from ppsync.config import Config
from ppsync.models import Status
from ppsync.values import display
from tools import make_test_workbook


def main() -> int:
    sandbox = Path(tempfile.mkdtemp(prefix="pp_sync_check_"))
    print(f"Песочница: {sandbox}")

    target_path = sandbox / "ПП_боевой.xlsx"
    source_path = sandbox / "ПП_копия.xlsx"
    make_test_workbook.build(target_path)
    cases = make_test_workbook.build_changed(target_path, source_path)
    print(f"Заложено различий: {len(cases)}")

    config = Config()
    config.target_file = str(target_path)
    config.source_file = str(source_path)
    config.data_dir = str(sandbox / "данные")
    config.sheet_name = make_test_workbook.SHEET
    config.ensure_dirs()

    session = SyncSession(config, progress=lambda text: print(f"  … {text}"))
    compare = session.compare()

    print("\nСтатусы:")
    for status, count in compare.counts().items():
        print(f"  {status.value:<22} {count}")

    applicable = [r for r in compare.rows if r.applicable]
    if len(applicable) < 2:
        print("!! Слишком мало применимых различий — проверка бессмысленна")
        return 2

    # После сравнения не должно быть отмечено ничего: пишем только то, что
    # человек отметил сам.
    if any(r.selected for r in compare.rows):
        print("!! После сравнения строки уже отмечены — так быть не должно")
        return 1

    clears = [r for r in applicable if r.status is Status.CLEAR]
    if not clears:
        print("!! В тестовых данных нет очистки — проверка неполная")
    for row in applicable:
        row.selected = True

    # Одну обычную правку намеренно оставляем неотмеченной — проверяем, что
    # неотмеченное не пишется. Очистку для этого не используем: её надо
    # проверить именно записью.
    skipped = next(r for r in applicable if r.status is Status.TO_APPLY)
    skipped.selected = False

    expected = {(r.address, r.new_display) for r in applicable if r.selected}
    result = session.apply(compare.rows)

    print(f"\nИтог: {result.summary()}")
    print(f"Бэкап  : {result.backup_path}")
    print(f"Журнал : {result.journal_path}")
    for error in result.errors:
        print(f"  ОШИБКА: {error}")
    for warning in result.verify_warnings:
        print(f"  ВНИМАНИЕ: {warning}")

    failures = check(target_path, config, expected, skipped, result, compare.rows)
    print()
    if failures:
        for failure in failures:
            print(f"ПРОВАЛ: {failure}")
        return 1
    print("Все проверки пройдены.")
    print(f"(песочница осталась: {sandbox})")
    return 0


def check(target_path, config, expected, skipped, result, rows) -> list[str]:
    """Независимая проверка результата через openpyxl."""
    failures: list[str] = []

    values = load_workbook(target_path, data_only=True)
    formulas = load_workbook(target_path)  # без data_only — видно формулы
    try:
        sheet = values[config.sheet_name]
        for address, expected_value in expected:
            actual = display(sheet[address].value)
            if actual != expected_value:
                failures.append(
                    f"{address}: в файле «{actual}», ожидалось «{expected_value}»"
                )
        if display(sheet[skipped.address].value) == skipped.new_display:
            failures.append(
                f"{skipped.address}: неотмеченное различие всё-таки записано"
            )

        formula_sheet = formulas[config.sheet_name]
        for row in rows:
            if row.status is Status.FORMULA and row.row:
                cell_value = str(formula_sheet[row.address].value or "")
                if not cell_value.startswith("="):
                    failures.append(f"{row.address}: формула затёрта ({cell_value!r})")
            if row.status is Status.KEY_NOT_FOUND and row.row:
                failures.append(f"{row.key}: у отсутствующей строки появился адрес")
    finally:
        values.close()
        formulas.close()

    if result.backup_path is None or not Path(result.backup_path).exists():
        failures.append("нет резервной копии")
    if result.journal_path is None or not Path(result.journal_path).exists():
        failures.append("нет журнала")
    return failures


if __name__ == "__main__":
    raise SystemExit(main())
