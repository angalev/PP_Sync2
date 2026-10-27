"""Проверка: боевой файл изменился между сравнением и записью.

    python -m tools.stale_check

Между «Сравнить» и «Применить» боевой файл никем не открыт, и коллеги могут
его менять. Здесь это разыгрывается на синтетической паре: после сравнения в
«боевой» файл вставляется строка, вставляется столбец и чужой рукой правится
одна из ячеек, которые человек отметил. Программа обязана:

* найти каждую строку заново по ключу X/Y/Z, а столбец — по коду, и записать
  по новым адресам (всё съехало на строку вниз и на столбец вправо);
* ячейку с чужой правкой не трогать и сказать об этом.
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from ppsync.applier import SyncSession
from ppsync.config import Config
from ppsync.models import Status
from ppsync.values import display
from tools import make_test_workbook

FOREIGN = "чужая правка"


def main() -> int:
    sandbox = Path(tempfile.mkdtemp(prefix="pp_sync_stale_"))
    print(f"Песочница: {sandbox}")
    target_path = sandbox / "ПП_боевой.xlsx"
    source_path = sandbox / "ПП_копия.xlsx"
    make_test_workbook.build(target_path)
    make_test_workbook.build_changed(target_path, source_path)

    config = Config()
    config.target_file = str(target_path)
    config.source_file = str(source_path)
    config.data_dir = str(sandbox / "данные")
    config.sheet_name = make_test_workbook.SHEET

    session = SyncSession(config, progress=lambda text: print(f"  … {text}"))
    compare = session.compare()
    applicable = [row for row in compare.rows if row.applicable]
    for row in applicable:
        row.selected = True
    before = {id(row): row.address for row in applicable}
    victim = next(row for row in applicable if row.status is Status.TO_APPLY)
    print(f"\nОтмечено {len(applicable)}; чужая правка ляжет в {victim.address}")

    # --- «коллега» меняет боевой файл, пока человек выбирает галочки ---
    workbook = load_workbook(target_path)
    sheet = workbook[config.sheet_name]
    sheet[victim.address] = FOREIGN
    sheet.insert_rows(config.first_data_row)  # все строки данных съезжают вниз
    sheet.insert_cols(2)  # все столбцы правее A съезжают вправо, вместе с кодами
    workbook.save(target_path)
    workbook.close()

    result = session.apply(compare.rows)
    print(f"\nИтог: {result.summary()}")
    for warning in result.verify_warnings:
        print(f"  ВНИМАНИЕ: {warning}")
    for error in result.errors:
        print(f"  ОШИБКА: {error}")

    failures: list[str] = []
    if result.stale != 1:
        failures.append(f"чужих правок замечено {result.stale}, а должна быть одна")
    if result.applied != len(applicable) - 1:
        failures.append(
            f"записано {result.applied}, а должно быть {len(applicable) - 1}"
        )
    if result.failed:
        failures.append(f"ошибок записи: {result.failed}")

    values = load_workbook(target_path, data_only=True)
    try:
        sheet = values[config.sheet_name]
        for row in applicable:
            moved = row.address != before[id(row)]
            actual = display(sheet[row.address].value) if moved else None
            if row is victim:
                # Адрес у неё не обновлялся: ячейку нашли, увидели чужое
                # значение и отступили. Ищем её по сдвигу вручную.
                shifted = sheet.cell(
                    row=row.row + 1, column=sheet[row.address].column + 1
                )
                if shifted.value != FOREIGN:
                    failures.append(f"чужая правка затёрта: там {shifted.value!r}")
                continue
            if not moved:
                failures.append(f"{row.key}: адрес не обновился ({row.address})")
            elif actual != row.new_display:
                failures.append(
                    f"{row.address}: в файле «{actual}», ожидалось «{row.new_display}»"
                )
            else:
                print(
                    f"  OK {before[id(row)]} -> {row.address}: «{actual}» "
                    f"({row.status.value})"
                )
    finally:
        values.close()

    print()
    if failures:
        for failure in failures:
            print(f"ПРОВАЛ: {failure}")
        return 1
    print("Все проверки пройдены.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
