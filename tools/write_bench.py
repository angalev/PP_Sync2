"""Замер скорости: сколько стоит внести сотню ячеек в боевой файл.

Отвечает на вопрос «почему на нескольких десятках ячеек всё встаёт». Два
прогона:

1. **Полный цикл** ``SyncSession`` на копии с ``--cells`` подсаженными
   правками — столько пользователь ждёт на самом деле, с разбивкой по фазам
   (открытие, карта ключей, чтение, сравнение, запись, save, сверка).
2. **Микрозамер стратегий записи** на той же книге: поячеечно с
   автопересчётом (как сейчас), поячеечно с ручным пересчётом, через сырой
   COM ``Value2``, и диапазонами по соседним ячейкам строки.

    python -m tools.write_bench --source "C:\\...\\Производственная программа 2026.xlsm"

Работает во временной песочнице. Исходный файл только читается.
"""

from __future__ import annotations

import argparse
import shutil
import sys
import tempfile
import threading
import time
from contextlib import suppress
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from ppsync.applier import SyncSession
from ppsync.config import Config
from ppsync.excel_guard import find_dialogs
from ppsync.models import Status
from tools.real_file_check import code_letters

XL_CALCULATION_AUTOMATIC = -4105
XL_CALCULATION_MANUAL = -4135


@dataclass
class Phases:
    """Отметки времени по сообщениям прогресса."""

    started: float = field(default_factory=time.perf_counter)
    marks: list[tuple[float, str]] = field(default_factory=list)

    def note(self, message: str) -> None:
        self.marks.append((time.perf_counter(), message))

    def report(self, title: str, finished: float) -> None:
        """Сообщение печатается ПЕРЕД работой, поэтому фаза — это интервал
        от своей отметки до следующей, а не до предыдущей."""
        print(f"\n{title}")
        # «Пишу N из M» — сотня одинаковых сообщений, сворачиваем в одну строку
        collapsed: list[tuple[float, str]] = []
        for stamp, message in self.marks:
            label = "запись ячеек" if message.startswith("Пишу ") else message
            if collapsed and collapsed[-1][1] == label:
                continue  # у пачки одинаковых берём отметку первого сообщения
            collapsed.append((stamp, label))
        for position, (stamp, label) in enumerate(collapsed):
            nxt = (
                collapsed[position + 1][0]
                if position + 1 < len(collapsed)
                else finished
            )
            print(f"  {nxt - stamp:7.2f} с  {label[:70]}")
        print(f"  {'-' * 7}")
        print(f"  {finished - self.started:7.2f} с  ИТОГО")


def unblock(path: Path) -> None:
    """Снять «файл скачан из интернета» — иначе Excel откроет в защищённом режиме."""
    zone = Path(str(path) + ":Zone.Identifier")
    with suppress(OSError):
        if zone.exists():
            zone.unlink()


def find_empty_cells(
    path: Path, sheet_name: str, count: int, skip: int = 0
) -> list[tuple[int, int]]:
    """Найти ``count`` пустых ячеек без формул в строках с ключом X/Y/Z.

    Возвращает пары ``(строка, столбец)``. По одной ячейке на строку — это
    худший случай для группировки в диапазоны и самый частый на практике:
    человек правит разные объекты, а не соседние колонки одной строки.
    """
    workbook = openpyxl.load_workbook(path, keep_vba=True, read_only=False)
    sheet = workbook[sheet_name]
    found: list[tuple[int, int]] = []
    skipped = 0
    # Программа сравнивает только столбцы с кодом в строке 9, а расчётные не
    # пишет — пустую ячейку ищем среди остальных, в правой половине листа.
    letters = code_letters(sheet)
    guarded = Config().protected_codes()
    key_column = column_index_from_string(letters[39])
    columns = sorted(
        column_index_from_string(letter)
        for code, letter in letters.items()
        if str(code) not in guarded
    )
    columns = [column for column in columns if column >= 60]
    try:
        for row in range(10, min(sheet.max_row, 1400) + 1):
            if not sheet.cell(row=row, column=key_column).value:
                continue
            for column in columns:
                cell = sheet.cell(row=row, column=column)
                if cell.value is not None:
                    continue
                if skipped < skip:
                    skipped += 1
                    break
                found.append((row, column))
                break
            if len(found) >= count:
                break
    finally:
        workbook.close()
    return found


def plant(path: Path, sheet_name: str, cells: list[tuple[int, int]]) -> None:
    """Записать в копию пометки — они и станут различиями «к применению»."""
    workbook = openpyxl.load_workbook(path, keep_vba=True)
    sheet = workbook[sheet_name]
    for number, (row, column) in enumerate(cells, start=1):
        sheet.cell(row=row, column=column).value = f"БЕНЧ-{number:03d}"
    workbook.save(path)
    workbook.close()


def guard(session: SyncSession, stop: threading.Event) -> None:
    """Сторож модальных окон: иначе зависший диалог не отличить от долгой работы."""
    reported: set[str] = set()
    while not stop.wait(5):
        titles = [t for pid in session.excel_pids for t in find_dialogs(pid)]
        for title in titles:
            if title not in reported:
                reported.add(title)
                print(f"  !! Excel показал окно «{title}» и ждёт ответа", flush=True)


def full_cycle(origin: Path, sandbox: Path, sheet: str, count: int) -> None:
    """Прогон 1: полный цикл на копии с ``count`` правками."""
    target_path = sandbox / "боевой.xlsm"
    source_path = sandbox / "копия_с_изменениями.xlsm"
    print(f"Копирую {origin.name} ({origin.stat().st_size / 1_048_576:.1f} МБ)…")
    shutil.copy2(origin, target_path)
    shutil.copy2(origin, source_path)
    unblock(target_path)
    unblock(source_path)

    print(f"Ищу {count} пустых ячеек без формул…")
    cells = find_empty_cells(source_path, sheet, count)
    print(f"  нашлось {len(cells)}, правлю копию…")
    plant(source_path, sheet, cells)

    config = Config()
    config.target_file = str(target_path)
    config.source_file = str(source_path)
    config.data_dir = str(sandbox / "данные")
    config.sheet_name = sheet
    config.ensure_dirs()

    phases = Phases()
    session = SyncSession(config, progress=phases.note)
    stop = threading.Event()
    threading.Thread(target=guard, args=(session, stop), daemon=True).start()

    try:
        compare = session.compare()
        planted = {f"{get_column_letter(c)}{r}" for r, c in cells}
        rows = [r for r in compare.rows if r.address in planted]
        applicable = [r for r in rows if r.applicable]
        print(
            f"\nРазличий всего: {len(compare.rows)}; "
            f"из подсаженных нашлось {len(rows)}, применимых {len(applicable)}"
        )
        other: dict[Status, int] = {}
        for row in compare.rows:
            other[row.status] = other.get(row.status, 0) + 1
        for status, number in sorted(other.items(), key=lambda kv: -kv[1]):
            print(f"  {status.value:<22} {number}")

        for row in compare.rows:
            row.selected = row.status is Status.TO_APPLY
        selected = [r for r in compare.rows if r.selected]
        print(f"\nПишу {len(selected)} ячеек…")
        result = session.apply(compare.rows)
        finished = time.perf_counter()
        print(f"Применение: {result.summary()}")
    finally:
        stop.set()

    phases.report(f"ФАЗЫ ПОЛНОГО ЦИКЛА ({count} ячеек)", finished)


def micro(origin: Path, sandbox: Path, sheet_name: str, count: int) -> None:
    """Прогон 2: сравнение стратегий записи на одной открытой книге."""
    import xlwings as xw

    path = sandbox / "микрозамер.xlsm"
    shutil.copy2(origin, path)
    unblock(path)

    print(f"\nПодбираю по {count} ячеек на каждую стратегию…")
    pool = find_empty_cells(path, sheet_name, count * 4)
    if len(pool) < count * 4:
        print(f"  !! нашлось только {len(pool)}, замер будет на меньшем числе")
    batches = [pool[i * count : (i + 1) * count] for i in range(4)]

    app = xw.App(visible=False, add_book=False)
    app.display_alerts = False
    app.screen_updating = False
    with suppress(Exception):
        app.api.AutomationSecurity = 3
    with suppress(Exception):
        app.api.EnableEvents = False
    with suppress(Exception):
        app.api.AskToUpdateLinks = False

    try:
        opened = time.perf_counter()
        book = app.books.open(str(path), update_links=False)
        sheet = book.sheets[sheet_name]
        print(f"  открытие книги: {time.perf_counter() - opened:.2f} с")

        formulas = int(
            app.api.WorksheetFunction.CountIf(sheet.api.UsedRange, "=*") or 0
        )
        print(f"  формул на листе (оценка): {formulas}")

        # Прогрев: первая запись строит дерево зависимостей и стоит дороже
        sheet.range((batches[0][0][0], batches[0][0][1])).value = "прогрев"

        results: list[tuple[str, float]] = []

        app.api.Calculation = XL_CALCULATION_AUTOMATIC
        started = time.perf_counter()
        for row, column in batches[0]:
            sheet.range((row, column)).value = "А"
        results.append(
            ("поячеечно, автопересчёт (как сейчас)", time.perf_counter() - started)
        )

        app.api.Calculation = XL_CALCULATION_MANUAL
        started = time.perf_counter()
        for row, column in batches[1]:
            sheet.range((row, column)).value = "Б"
        results.append(("поячеечно, ручной пересчёт", time.perf_counter() - started))

        started = time.perf_counter()
        api = sheet.api
        for row, column in batches[2]:
            api.Cells(row, column).Value2 = "В"
        results.append(
            ("сырой COM Value2, ручной пересчёт", time.perf_counter() - started)
        )

        # Соседние ячейки одной строки — лучший случай для группировки
        row_number = batches[3][0][0]
        first_column = batches[3][0][1]
        started = time.perf_counter()
        sheet.range(
            (row_number, first_column), (row_number, first_column + count - 1)
        ).value = [["Г"] * count]
        results.append(
            (f"одним диапазоном {count} ячеек в строке", time.perf_counter() - started)
        )

        started = time.perf_counter()
        app.api.Calculate()
        recalc = time.perf_counter() - started

        started = time.perf_counter()
        book.save()
        save = time.perf_counter() - started

        print(f"\nСТРАТЕГИИ ЗАПИСИ ({count} ячеек каждая)")
        base = results[0][1]
        for name, seconds in results:
            speedup = f"×{base / seconds:.0f} быстрее" if seconds > 0 else ""
            per_cell = seconds / max(count, 1) * 1000
            print(f"  {seconds:8.2f} с  ({per_cell:6.1f} мс/ячейка)  {name}  {speedup}")
        print(f"\n  {recalc:8.2f} с  один полный пересчёт книги")
        print(f"  {save:8.2f} с  сохранение")
    finally:
        with suppress(Exception):
            app.api.Calculation = XL_CALCULATION_AUTOMATIC
        with suppress(Exception):
            book.close()
        with suppress(Exception):
            app.quit()
        with suppress(Exception):
            app.kill()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, help="боевой файл-образец")
    parser.add_argument("--sheet", default="ПП 2026")
    parser.add_argument("--cells", type=int, default=100)
    parser.add_argument("--skip-full", action="store_true")
    parser.add_argument("--skip-micro", action="store_true")
    args = parser.parse_args()

    origin = Path(args.source)
    if not origin.exists():
        print(f"Не найден файл: {origin}")
        return 2

    sandbox = Path(tempfile.mkdtemp(prefix="pp_write_bench_"))
    try:
        if not args.skip_full:
            full_cycle(origin, sandbox, args.sheet, args.cells)
        if not args.skip_micro:
            micro(origin, sandbox, args.sheet, args.cells)
    finally:
        print(f"\nПесочница: {sandbox}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
