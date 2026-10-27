r"""Полный перенос из копии в боевой файл и детальная сверка результата.

    python -m tools.transfer_check --target "...\боевой.xlsm" --source "...\копия.xlsm"

Работает на копиях в песочнице, исходные файлы только читаются. Файлы могут
быть разной ширины (старый формат ПП и 2.0): столбцы сопоставляются по кодам
из строки 9, как и в самой программе.

Что делает:

1. сравнивает боевой файл с копией и отмечает **всё**, что можно применить;
2. записывает, сохраняет, закрывает Excel;
3. открывает записанный файл заново — уже через openpyxl, независимо от
   программы — и сверяет его с копией ячейка за ячейкой по всей области
   сравнения, а не только по применённым правкам;
4. раскладывает каждое оставшееся расхождение по причине, чтобы было видно,
   где программа не перенесла осознанно, а где недосмотрела.

Отдельно сходится число формул: сколько было в боевом, сколько стало и сколько
ленивых ссылок заменено значением.
"""

from __future__ import annotations

import argparse
import collections
import datetime as _dt
import shutil
import sys
import tempfile
import time
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from ppsync.applier import SyncSession
from ppsync.config import Config
from ppsync.values import display, is_empty_formula_result, looks_like_key

SAMPLE = 10  # сколько примеров печатать в каждой группе

#: Причины, из-за которых значение осталось непереносенным по замыслу.
EXPECTED_REASONS = (
    "формула с обеих сторон — сравнивать нечего",
    "формула в расчётном столбце боевого — не трогаем",
    "формула копии посчитала пусто (=ссылка на пустую ячейку)",
    "формула копии посчитала ошибку (#REF!, #DIV/0!…)",
)

_ERRORS = ("#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Перенос и детальная сверка")
    parser.add_argument("--target", required=True, help="боевой файл")
    parser.add_argument("--source", required=True, help="файл с изменениями")
    parser.add_argument("--sheet", default="ПП 2026")
    parser.add_argument("--keep", action="store_true", help="не удалять песочницу")
    return parser.parse_args()


def read_sheet(path: Path, sheet: str, *, formulas: bool) -> list[tuple[Any, ...]]:
    book = openpyxl.load_workbook(
        path, data_only=not formulas, read_only=True, keep_links=False
    )
    try:
        return list(book[sheet].iter_rows(values_only=True))
    finally:
        book.close()


def at(row: tuple[Any, ...], column: int) -> Any:
    return row[column - 1] if column - 1 < len(row) else None


def is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def count_formulas(grid: list[tuple[Any, ...]]) -> int:
    return sum(1 for row in grid for value in row if is_formula(value))


def code_columns(grid: list[tuple[Any, ...]], config: Config) -> dict[str, int]:
    """Строка кодов: ``код -> номер столбца`` (независимо от программы)."""
    columns: dict[str, int] = {}
    for position, value in enumerate(grid[config.code_row - 1], start=1):
        if isinstance(value, _dt.datetime):
            # Код в ячейке с форматом даты: 21 читается как 21.01.1900
            value = value.timetuple().tm_yday
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        if value not in (None, ""):
            columns.setdefault(str(value).strip(), position)
    return columns


def keymap(
    grid: list[tuple[Any, ...]], config: Config, key_column: int
) -> dict[str, int]:
    found: dict[str, int] = {}
    for number, row in enumerate(grid, start=1):
        if number < config.first_data_row:
            continue
        key = at(row, key_column)
        if looks_like_key(key):
            found.setdefault(display(key), number)
    return found


def main() -> int:
    args = parse_args()
    workdir = Path(tempfile.mkdtemp(prefix="pp_transfer_"))
    target = workdir / "боевой.xlsm"
    source = workdir / "копия.xlsm"
    shutil.copy2(args.target, target)
    shutil.copy2(args.source, source)
    print(f"Песочница: {workdir}")

    config = Config()
    config.sheet_name = args.sheet
    config.target_file = str(target)
    config.source_file = str(source)
    config.data_dir = str(workdir / "данные")

    formulas_before = count_formulas(read_sheet(target, args.sheet, formulas=True))

    session = SyncSession(config, progress=lambda text: print(f"  … {text}"))
    started = time.perf_counter()
    compare = session.compare()
    rows = compare.rows

    print()
    print(f"Сравнение за {time.perf_counter() - started:.1f} с")
    print(
        f"Строк: в боевом {compare.target.count}, в копии {compare.source.count}; "
        f"столбцов в сравнении: {len(compare.codes)}"
    )
    for warning in compare.warnings:
        print(f"  ! {warning}")
    print(f"Различий найдено: {len(rows)}")
    for status, count in collections.Counter(
        row.status.value for row in rows
    ).most_common():
        print(f"  {status:<22} {count}")

    print()
    print("Что переносим, по столбцам:")
    by_field = collections.Counter(
        (row.column, row.field_name) for row in rows if row.applicable
    )
    for (letter, name), count in by_field.most_common(15):
        print(f"  {letter:<4} {count:>5}  {name.replace(chr(10), ' ')[:58]}")

    replacing = [row for row in rows if row.applicable and row.replaces_formula]
    print()
    print(f"Из них заменят формулу в боевом: {len(replacing)}")
    for row in replacing[:SAMPLE]:
        print(f"  {row.address:<9} {row.current_display!r} -> {row.new_display!r}")

    code = 0
    applicable = [row for row in rows if row.applicable]
    if applicable:
        for row in applicable:
            row.selected = True
        started = time.perf_counter()
        result = session.apply(rows)
        print()
        print(
            f"Применение за {time.perf_counter() - started:.1f} с: {result.summary()}"
        )
        for text in result.errors[:SAMPLE]:
            print(f"  ошибка: {text}")
        for text in result.verify_warnings[:SAMPLE]:
            print(f"  внимание: {text}")
        if result.failed or result.stale or result.verify_warnings:
            code = 1
    else:
        print()
        print("Применять нечего.")

    code = max(
        code, verify(args, target, source, compare.codes, formulas_before, replacing)
    )
    if args.keep:
        print(f"Песочница оставлена: {workdir}")
    else:
        shutil.rmtree(workdir, ignore_errors=True)
    return code


def verify(
    args: argparse.Namespace,
    target: Path,
    source: Path,
    codes: list[str],
    formulas_before: int,
    replacing: list[Any],
) -> int:
    """Сверить записанный боевой файл с копией по всей области сравнения."""
    print()
    print("СВЕРКА ПОСЛЕ ЗАПИСИ")

    config = Config()
    protected = config.protected_codes()
    target_formulas = read_sheet(target, args.sheet, formulas=True)
    source_formulas = read_sheet(source, args.sheet, formulas=True)
    target_values = read_sheet(target, args.sheet, formulas=False)
    source_values = read_sheet(source, args.sheet, formulas=False)

    formulas_after = count_formulas(target_formulas)
    expected = formulas_before - len(replacing)
    print(f"  формул в боевом: было {formulas_before}, стало {formulas_after}")
    print(f"  ожидалось {expected} (минус {len(replacing)} заменённых ссылок)")

    target_columns = code_columns(target_formulas, config)
    source_columns = code_columns(source_formulas, config)
    target_rows = keymap(target_values, config, target_columns[config.key_code])
    source_rows = keymap(source_values, config, source_columns[config.key_code])
    missing_keys = [key for key in source_rows if key not in target_rows]

    reasons: collections.Counter[str] = collections.Counter()
    examples: dict[str, list[str]] = collections.defaultdict(list)
    compared = 0
    for key, source_row in source_rows.items():
        target_row = target_rows.get(key)
        if target_row is None:
            continue
        for code in codes:
            target_column = target_columns[code]
            source_column = source_columns[code]
            compared += 1
            left = display(at(target_values[target_row - 1], target_column))
            right = display(at(source_values[source_row - 1], source_column))
            if left == right:
                continue
            in_source = is_formula(at(source_formulas[source_row - 1], source_column))
            in_target = is_formula(at(target_formulas[target_row - 1], target_column))
            source_value = at(source_values[source_row - 1], source_column)
            if in_source and is_empty_formula_result(source_value):
                # Ноль по пустой ссылке: в ячейке с форматом даты это «00:00».
                reason = EXPECTED_REASONS[2]
            elif source_value in _ERRORS:
                reason = EXPECTED_REASONS[3]
            elif in_source and in_target:
                reason = EXPECTED_REASONS[0]
            elif in_target and code in protected:
                reason = EXPECTED_REASONS[1]
            elif in_target:
                reason = "ОСТАЛАСЬ формула в обычном столбце"
            elif in_source:
                reason = "НЕ ПЕРЕНЕСЛОСЬ значение формулы копии"
            else:
                reason = "НЕ ПЕРЕНЕСЛОСЬ обычное значение"
            reasons[reason] += 1
            if len(examples[reason]) < SAMPLE:
                examples[reason].append(
                    f"код {code}, {get_column_letter(target_column)}{target_row}: "
                    f"в боевом {left!r}, в копии {right!r}"
                )

    mismatched = sum(reasons.values())
    print(f"  сверено ячеек: {compared}, совпало: {compared - mismatched}")
    if missing_keys:
        print(f"  кодов X/Y/Z из копии нет в боевом: {len(missing_keys)}")

    print()
    problems = 0
    for reason, count in reasons.most_common():
        print(f"  {count:>6}  {reason}")
        for line in examples[reason]:
            print(f"          {line}")
        if reason not in EXPECTED_REASONS:
            problems += count

    print()
    if formulas_after != expected:
        print(f"!! формулы не сошлись: ждали {expected}, получили {formulas_after}")
        problems += 1
    if problems:
        print(f"НЕ ВСЁ: {problems} расхождений требуют разбора.")
        return 1
    print("Перенос полный: всё, что можно было перенести, перенесено.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
