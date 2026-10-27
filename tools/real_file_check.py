"""Проверка на копии боевого файла: что уцелело после записи через xlwings.

Из боевого файла делаются две копии — «боевая» и «копия с изменениями», во
второй правится несколько ячеек (пустая с датным форматом, числовая,
заполненная, с формулой, с текстовым форматом). Затем прогоняется полный цикл
и сравниваются слепки до/после: формулы всех листов, оформление ячеек, состав
частей книги, объединения, условное форматирование, проверки данных, имена.

    python -m tools.real_file_check --source "C:\\...\\Производственная программа 2026.xlsm"

Работает во временной песочнице. Исходный файл только читается.
"""

from __future__ import annotations

import argparse
import datetime as _dt
import random
import shutil
import sys
import tempfile
import threading
import time
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from ppsync.applier import SyncSession
from ppsync.config import Config
from ppsync.excel_guard import find_dialogs
from ppsync.models import Status

SAMPLE_CELLS = 400  # сколько случайных ячеек проверять на форматирование


@dataclass
class Snapshot:
    """Слепок книги для сравнения «до/после»."""

    zip_entries: dict[str, int] = field(default_factory=dict)
    formulas: dict[str, str] = field(default_factory=dict)
    styles: dict[str, str] = field(default_factory=dict)
    merged: dict[str, int] = field(default_factory=dict)
    conditional: dict[str, int] = field(default_factory=dict)
    validations: dict[str, int] = field(default_factory=dict)
    sheet_names: list[str] = field(default_factory=list)
    defined_names: list[str] = field(default_factory=list)
    values: dict[str, str] = field(default_factory=dict)


def style_key(cell: Any) -> str:
    """Строка, описывающая оформление ячейки."""
    font, fill, border, align = cell.font, cell.fill, cell.border, cell.alignment
    sides = "|".join(
        f"{name}:{getattr(border, name).style or '-'}"
        for name in ("left", "right", "top", "bottom")
    )
    return ";".join(
        [
            f"fmt={cell.number_format}",
            f"font={font.name},{font.sz},b={font.b},i={font.i},color={_rgb(font.color)}",
            f"fill={fill.patternType},{_rgb(fill.fgColor)},{_rgb(fill.bgColor)}",
            f"border={sides}",
            f"align={align.horizontal},{align.vertical},wrap={align.wrap_text}",
        ]
    )


def _rgb(color: Any) -> str:
    if color is None:
        return "-"
    return str(getattr(color, "rgb", None) or getattr(color, "theme", None) or "-")


def snapshot(path: Path, sample: list[str], sheet_name: str) -> Snapshot:
    """Снять слепок: zip-состав, формулы, оформление, структуры листов."""
    shot = Snapshot()

    with zipfile.ZipFile(path) as archive:
        for info in archive.infolist():
            shot.zip_entries[info.filename] = info.CRC

    workbook = openpyxl.load_workbook(path, keep_vba=True)
    try:
        shot.sheet_names = list(workbook.sheetnames)
        shot.defined_names = sorted(
            list(workbook.defined_names.keys())
            + [
                f"{sheet.title}!{name}"
                for sheet in workbook.worksheets
                for name in sheet.defined_names
            ]
        )
        for name in workbook.sheetnames:
            sheet = workbook[name]
            shot.merged[name] = len(sheet.merged_cells.ranges)
            shot.conditional[name] = len(sheet.conditional_formatting._cf_rules)
            shot.validations[name] = len(sheet.data_validations.dataValidation)
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        shot.formulas[f"{name}!{cell.coordinate}"] = cell.value

        target = workbook[sheet_name]
        for address in sample:
            cell = target[address]
            shot.styles[address] = style_key(cell)
            shot.values[address] = repr(cell.value)
    finally:
        workbook.close()
    return shot


def make_changes(source: Path, sheet_name: str) -> tuple[list[str], dict[str, str]]:
    """Внести в копию правки разного типа.

    Возвращает (адреса под наблюдением, пояснения по случаям).
    """
    workbook = openpyxl.load_workbook(source, keep_vba=True)
    sheet = workbook[sheet_name]

    cases: dict[str, str] = {}
    watch: list[str] = []

    # Excel хранит даты без таймзоны — naive datetime здесь намеренный
    test_date = _dt.datetime(2026, 6, 30)  # noqa: DTZ001

    # Столбцы заданы кодами из строки 9 формы от 14.10.2026: 264 «ФАКТ, дата
    # сдачи», 102 «Ожидаемое_дек», 265 «Номер ЗЭПБ» (текстовый формат).
    plan = [
        (
            264,
            test_date,
            "дата в пустую ячейку с форматом даты",
            lambda c: c.value is None,
        ),
        (
            102,
            1,
            "число в пустую ячейку с дробным форматом",
            lambda c: c.value is None,
        ),
        (
            316,
            "ТА-ПРОВЕРКА-001",
            "текст поверх заполненной ячейки (перезапись)",
            lambda c: c.value is not None and not str(c.value).startswith("="),
        ),
        (
            313,
            "не должно записаться",
            "значение поверх формулы (ожидаем блокировку)",
            lambda c: isinstance(c.value, str) and c.value.startswith("="),
        ),
        (
            265,
            "15.07.2026",
            "дата в пустую ячейку с текстовым форматом",
            lambda c: c.value is None,
        ),
        (
            312,
            777,
            "значение поверх ленивой ссылки (формула заменяется)",
            lambda c: isinstance(c.value, str) and c.value.startswith("="),
        ),
    ]

    letters = code_letters(sheet)
    key_letter = letters[39]
    used_rows: set[int] = set()
    for code, new_value, case, predicate in plan:
        column = letters[code]
        for row in range(10, min(sheet.max_row, 1300) + 1):
            if row in used_rows or not sheet[f"{key_letter}{row}"].value:
                continue
            cell = sheet[f"{column}{row}"]
            if not predicate(cell):
                continue
            used_rows.add(row)
            address = f"{column}{row}"
            cell.value = new_value
            cases[address] = case
            watch.append(address)
            watch.append(f"{get_column_letter(cell.column - 1)}{row}")
            watch.append(f"{get_column_letter(cell.column + 1)}{row}")
            watch.append(f"{column}{row - 1}")
            watch.append(f"{column}{row + 1}")
            break

    workbook.save(source)
    workbook.close()

    # плюс случайная выборка по листу — вдруг оформление поедет где-то ещё
    rng = random.Random(20260814)
    width = len(letters)
    for _ in range(SAMPLE_CELLS):
        watch.append(
            f"{get_column_letter(rng.randint(1, width))}{rng.randint(1, 1300)}"
        )
    return sorted(set(watch)), cases


def code_letters(sheet: Any, code_row: int = 9) -> dict[int, str]:
    """Строка кодов листа: ``код -> буква столбца``.

    Ячейка с кодом может стоять в формате даты (в ПП 2.0 коды 21 и 22), и
    тогда openpyxl отдаёт дату января 1900 года — код в ней номер дня.
    """
    letters: dict[int, str] = {}
    for cell in sheet[code_row]:
        value = cell.value
        if isinstance(value, _dt.datetime):
            value = value.timetuple().tm_yday
        if isinstance(value, (int, float)):
            letters[int(value)] = cell.column_letter
    return letters


def unblock(path: Path) -> None:
    """Снять «файл скачан из интернета» — иначе Excel откроет в защищённом режиме."""
    zone = Path(str(path) + ":Zone.Identifier")
    try:
        if zone.exists():
            zone.unlink()
    except OSError:
        pass


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, help="боевой файл-образец")
    parser.add_argument("--sheet", default="ПП 2026")
    args = parser.parse_args()

    origin = Path(args.source)
    if not origin.exists():
        print(f"Не найден файл: {origin}")
        return 2

    sandbox = Path(tempfile.mkdtemp(prefix="pp_real_check_"))
    print(f"Копирую {origin.name} ({origin.stat().st_size / 1_048_576:.1f} МБ)…")
    target_path = sandbox / "боевой.xlsm"
    source_path = sandbox / "копия_с_изменениями.xlsm"
    shutil.copy2(origin, target_path)
    shutil.copy2(origin, source_path)
    unblock(target_path)
    unblock(source_path)

    print("Правлю копию…")
    watch, cases = make_changes(source_path, args.sheet)
    for address, case in cases.items():
        print(f"  {address:<8} {case}")

    print("Снимаю слепок ДО…")
    before = snapshot(target_path, watch, args.sheet)
    print(
        f"  формул в книге: {len(before.formulas)}; ячеек под наблюдением: {len(watch)}"
    )

    config = Config()
    config.target_file = str(target_path)
    config.source_file = str(source_path)
    config.data_dir = str(sandbox / "данные")
    config.sheet_name = args.sheet
    config.ensure_dirs()

    session = SyncSession(
        config, progress=lambda text: print(f"  … {text}", flush=True)
    )

    # Сторож модальных окон: без него зависший диалог Excel не отличить от
    # долгой работы, и прогон стоит часами.
    stop_watch = threading.Event()

    def watch_dialogs() -> None:
        reported: set[str] = set()
        while not stop_watch.wait(5):
            titles = [t for pid in session.excel_pids for t in find_dialogs(pid)]
            for title in titles:
                if title not in reported:
                    reported.add(title)
                    print(
                        f"  !! Excel показал окно «{title}» и ждёт ответа", flush=True
                    )

    threading.Thread(target=watch_dialogs, daemon=True).start()

    started = time.perf_counter()
    compare = session.compare()
    print(f"\nСравнение построено за {time.perf_counter() - started:.1f} с")
    print(
        f"Ключей X/Y/Z: {compare.target.count}, "
        f"столбцов в сравнении: {len(compare.codes)}"
    )
    print(f"Различий: {len(compare.rows)}")
    for warning in compare.warnings:
        print(f"  ! {warning}")

    print("\nЧто нашлось в наших целевых ячейках:")
    planted = [r for r in compare.rows if r.address in cases]
    for row in planted:
        print(
            f"  {row.address:<8} {row.status.value:<20} "
            f"было={row.current_display or '—'!r:<22} станет={row.new_display!r}"
        )
    if len(planted) < len(cases):
        missing = set(cases) - {r.address for r in planted}
        print(f"  !! не найдены как различия: {sorted(missing)}")

    # Программа сама не отмечает ничего — пишется только то, что отметил
    # человек. В автопрогоне «человек» отмечает всё применимое.
    if any(r.selected for r in compare.rows):
        print("  !! после сравнения строки уже отмечены — так быть не должно")
    for row in compare.rows:
        row.selected = row.applicable

    result = session.apply(compare.rows)
    stop_watch.set()
    print(f"\nПрименение: {result.summary()}")
    for error in result.errors:
        print(f"  ОШИБКА: {error}")
    for warning in result.verify_warnings:
        print(f"  ВНИМАНИЕ: {warning}")

    print("\nСнимаю слепок ПОСЛЕ…")
    after = snapshot(target_path, watch, args.sheet)

    failures = compare_snapshots(before, after, compare.rows, cases, args.sheet)
    print()
    if failures:
        print("НАЙДЕНЫ РАСХОЖДЕНИЯ:")
        for failure in failures:
            print(f"  • {failure}")
    else:
        print("Расхождений нет: формулы, оформление и служебные части книги на месте.")

    print(f"\nПесочница: {sandbox}")
    return 1 if failures else 0


def compare_snapshots(
    before: Snapshot,
    after: Snapshot,
    rows: list,
    cases: dict,
    sheet_name: str = "ПП 2026",
) -> list[str]:
    """Сравнить слепки. Ожидаемые изменения (записанные ячейки) не считаются."""
    failures: list[str] = []
    written = {r.address for r in rows if r.selected and r.applicable}

    # --- состав архива ---
    missing = set(before.zip_entries) - set(after.zip_entries)
    added = set(after.zip_entries) - set(before.zip_entries)
    critical = [
        n
        for n in missing
        if any(
            key in n
            for key in (
                "vbaProject",
                "pivotCache",
                "pivotTable",
                "ctrlProps",
                "drawing",
                "chart",
            )
        )
    ]
    if critical:
        failures.append(f"пропали части книги: {', '.join(sorted(critical))}")
    print(
        f"  части архива: было {len(before.zip_entries)}, стало {len(after.zip_entries)}"
    )
    if missing:
        print(f"    исчезли ({len(missing)}): {', '.join(sorted(missing)[:6])}")
    if added:
        print(f"    появились ({len(added)}): {', '.join(sorted(added)[:6])}")

    # --- формулы ---
    # Ленивая ссылка в обычном столбце заменяется значением по замыслу —
    # такая формула «потеряна» законно, остальные нет.
    replaced = {
        f"{sheet_name}!{r.address}"
        for r in rows
        if r.selected and r.applicable and r.replaces_formula
    }
    lost = {
        k: v
        for k, v in before.formulas.items()
        if k not in after.formulas and k not in replaced
    }
    kept = [k for k in replaced if k in after.formulas]
    if kept:
        failures.append(f"формула должна была замениться значением: {kept}")
    print(f"  заменено формул значением по замыслу: {len(replaced)}")
    changed = {
        k: (v, after.formulas[k])
        for k, v in before.formulas.items()
        if k in after.formulas and after.formulas[k] != v
    }
    new = {k: v for k, v in after.formulas.items() if k not in before.formulas}
    print(f"  формулы: было {len(before.formulas)}, стало {len(after.formulas)}")
    for key in list(lost)[:5]:
        print(f"    потеряна: {key} = {lost[key][:60]}")
    for key in list(changed)[:5]:
        print(f"    изменена: {key}: {changed[key][0][:40]} -> {changed[key][1][:40]}")
    for key in list(new)[:5]:
        print(f"    новая: {key} = {new[key][:60]}")
    if lost:
        failures.append(f"потеряно формул: {len(lost)}")
    if changed:
        failures.append(f"изменено формул: {len(changed)}")
    if new:
        failures.append(f"появилось формул: {len(new)}")

    # --- оформление ---
    style_changes = {
        address: (before.styles[address], after.styles[address])
        for address in before.styles
        if after.styles.get(address) != before.styles[address]
    }
    print(
        f"  оформление: проверено ячеек {len(before.styles)}, изменилось {len(style_changes)}"
    )
    for address, (was, now) in list(style_changes.items())[:10]:
        note = " (в неё писали)" if address in written else ""
        print(f"    {address}{note}: {_style_diff(was, now)}")
        failures.append(f"изменилось оформление {address}{note}")

    # --- значения ---
    print("\n  что стало с целевыми ячейками:")
    for row in rows:
        if row.address not in cases:
            continue
        print(
            f"    {row.address:<8} {row.status.value:<20} "
            f"{before.values.get(row.address, '?')} -> "
            f"{after.values.get(row.address, '?')}   {cases[row.address]}"
        )
        if row.status is Status.FORMULA and before.values.get(
            row.address
        ) != after.values.get(row.address):
            failures.append(
                f"{row.address}: формула затёрта, хотя строка была заблокирована"
            )

    neighbours = [
        address
        for address in before.values
        if address not in written
        and before.values[address] != after.values.get(address)
    ]
    if neighbours:
        failures.append(
            f"изменились значения ячеек, в которые не писали: {neighbours[:10]}"
        )

    # --- прочие структуры ---
    for name, title in (
        ("merged", "объединённых диапазонов"),
        ("conditional", "правил условного форматирования"),
        ("validations", "проверок данных"),
    ):
        was_value = getattr(before, name)
        now_value = getattr(after, name)
        if was_value != now_value:
            failures.append(
                f"изменилось количество {title}: {was_value} -> {now_value}"
            )
    if before.sheet_names != after.sheet_names:
        failures.append(
            f"изменился состав листов: {before.sheet_names} -> {after.sheet_names}"
        )
    if before.defined_names != after.defined_names:
        only_before = set(before.defined_names) - set(after.defined_names)
        only_after = set(after.defined_names) - set(before.defined_names)
        # Excel сам нормализует служебные имена автофильтра и представлений
        # (_FilterDatabase, *.wvu.*) при любом сохранении — проверено отдельно
        # прогоном «открыл и сохранил, ничего не записывая». Это не наши следы.
        churn = sorted(
            name
            for name in only_before | only_after
            if "_FilterDatabase" not in name and ".wvu." not in name
        )
        print(
            f"  служебные имена Excel: -{len(only_before)} +{len(only_after)} "
            "(нормализация автофильтра и представлений)"
        )
        if churn:
            failures.append(f"изменились именованные диапазоны: {churn}")

    return failures


def _style_diff(was: str, now: str) -> str:
    parts_was = dict(p.split("=", 1) for p in was.split(";"))
    parts_now = dict(p.split("=", 1) for p in now.split(";"))
    diffs = [
        f"{key}: {parts_was[key]} -> {parts_now[key]}"
        for key in parts_was
        if parts_was[key] != parts_now.get(key)
    ]
    return "; ".join(diffs) or "нет различий"


if __name__ == "__main__":
    raise SystemExit(main())
