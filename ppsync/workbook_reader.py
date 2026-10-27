"""Чтение листа ПП через xlwings (реальный Excel по COM).

И боевой файл, и копия с изменениями читаются одним и тем же кодом — поэтому
значения с обеих сторон приходят в одинаковом виде, и сравнивать их можно
напрямую.

Два правила модуля:

* **не читать по одной ячейке.** Каждое обращение к COM стоит миллисекунды,
  на 1300 строках поячеечное чтение превращается в минуты. Лист читается
  прямоугольником за один вызов: значения и формулы — по вызову на каждое;
* **столбцы находятся по коду, а не по букве.** Код стоит в строке кодов
  (в ПП это строка 9) и остаётся при своём столбце, сколько бы столбцов ни
  удалили и ни вставили слева. Буква нужна только чтобы показать адрес.
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from typing import Any

from openpyxl.utils import get_column_letter

from .config import Config
from .values import clean_str, code_text, display, from_excel, looks_like_key

#: Предохранитель: сколько строк максимум сканируем ниже первой строки данных.
#: UsedRange у живой книги иногда раздут пустым форматированием.
MAX_SCAN_ROWS = 50_000


class WorkbookStructureError(Exception):
    """Структура листа не соответствует настройкам. ``problems`` — причины."""

    def __init__(self, message: str, problems: list[str] | None = None) -> None:
        super().__init__(message)
        self.problems = problems or []

    def full_text(self) -> str:
        if not self.problems:
            return str(self)
        return str(self) + "\n\n" + "\n".join(f"• {p}" for p in self.problems)


@dataclass
class RowData:
    """Строка данных листа: значения и формулы всех столбцов с кодом."""

    row: int
    key: str
    transgaz: str
    object_name: str
    values: list[Any]  # позиция = номер столбца - 1
    formulas: dict[int, str] = field(default_factory=dict)  # столбец -> формула

    def value(self, column: int) -> Any:
        position = column - 1
        return self.values[position] if 0 <= position < len(self.values) else None

    def formula(self, column: int) -> str:
        """Текст формулы в ячейке; пустая строка — формулы нет."""
        return self.formulas.get(column, "")


@dataclass
class SheetData:
    """Прочитанный лист: строки по ключу X/Y/Z, столбцы по коду."""

    title: str  # «боевой файл» или «файл с изменениями» — для сообщений
    columns: dict[str, int] = field(default_factory=dict)  # код -> номер столбца
    headers: dict[str, str] = field(default_factory=dict)  # код -> имя столбца
    rows: dict[str, RowData] = field(default_factory=dict)
    duplicates: dict[str, list[int]] = field(default_factory=dict)
    # код -> числовой формат столбца, если он один на весь столбец (иначе "")
    formats: dict[str, str] = field(default_factory=dict)
    last_row: int = 0  # последняя строка, где реально нашёлся ключ

    @property
    def count(self) -> int:
        return len(self.rows)

    @property
    def codes(self) -> list[str]:
        """Коды столбцов слева направо."""
        return sorted(self.columns, key=lambda code: self.columns[code])

    def letter(self, code: str) -> str:
        column = self.columns.get(code)
        return get_column_letter(column) if column else ""

    def describe(self, code: str) -> str:
        """«316 (CN, Номер техакта)» — столбец для сообщений человеку."""
        name = self.headers.get(code, "")
        tail = ", ".join(part for part in (self.letter(code), name[:40]) if part)
        return f"{code} ({tail})" if tail else code


def read_sheet(
    sheet: Any, config: Config, title: str, *, with_formats: bool = False
) -> SheetData:
    """Прочитать лист целиком: коды, шапку, значения и формулы.

    ``sheet`` — лист xlwings. Границы берутся из ``UsedRange``, а НЕ через
    ``End(xlUp)``/``End(xlToLeft)``: на боевом файле почти всегда включён
    автофильтр (например, по одному трансгазу), а навигация End пропускает
    скрытые строки и столбцы — карта ключей тогда молча обрывается на
    последней видимой строке.

    ``with_formats`` — заодно прочитать числовые форматы столбцов (нужно
    только боевому файлу: по формату видно ячейки «Текстовый»).
    """
    data = SheetData(title=title)
    last_cell = sheet.used_range.last_cell
    last_row = min(int(last_cell.row), config.first_data_row + MAX_SCAN_ROWS)
    last_column = int(last_cell.column)

    data.columns = _read_codes(sheet, config, last_column, title)
    width = max(data.columns.values())
    data.headers = _read_headers(sheet, config, data.columns, width)

    key_column = data.columns.get(config.key_code)
    if key_column is None:
        raise WorkbookStructureError(
            f"В листе «{config.sheet_name}» ({title}) нет столбца с ключом X/Y/Z",
            [
                f"В строке {config.code_row} не найден код {config.key_code}.",
                (
                    "Проверьте настройки «строка кодов столбцов» и «код столбца "
                    "ключа X/Y/Z»."
                ),
                f"Коды, которые есть в строке: {', '.join(data.codes[:15])}…",
            ],
        )
    transgaz_column = data.columns.get(config.transgaz_code, 0)
    object_column = data.columns.get(config.object_code, 0)

    first = config.first_data_row
    if last_row >= first:
        block = sheet.range((first, 1), (last_row, width))
        rows = last_row - first + 1
        values = _as_rows(block.value, rows, width)
        formulas = _as_rows(block.formula, rows, width)
        for offset, row_values in enumerate(values):
            key = display(_at(row_values, key_column))
            # Одно и то же правило для боевого файла и для копии: иначе один
            # заводит ключ там, где другой строку пропускает, и строка молча
            # выпадает из сравнения навсегда.
            if not looks_like_key(key):
                continue
            row_number = first + offset
            data.last_row = row_number
            existing = data.rows.get(key)
            if existing is not None:
                data.duplicates.setdefault(key, [existing.row]).append(row_number)
                continue
            data.rows[key] = RowData(
                row=row_number,
                key=key,
                transgaz=display(_at(row_values, transgaz_column)),
                object_name=display(_at(row_values, object_column)),
                values=[from_excel(value) for value in row_values],
                formulas={
                    position: text
                    for position, text in enumerate(formulas[offset], start=1)
                    if isinstance(text, str) and text.startswith("=")
                },
            )

    if not data.rows:
        raise WorkbookStructureError(
            f"В листе «{config.sheet_name}» ({title}) не найдено ни одного ключа X/Y/Z",
            [
                (
                    f"Ключи искались в столбце {get_column_letter(key_column)} "
                    f"(код {config.key_code}), строки {first}–{last_row}."
                ),
                (
                    "Проверьте настройки «первая строка данных» и «код столбца "
                    "ключа X/Y/Z»."
                ),
            ],
        )
    if with_formats:
        data.formats = _read_column_formats(sheet, data.columns, first, data.last_row)
    return data


def _read_column_formats(
    sheet: Any, columns: dict[str, int], first_row: int, last_row: int
) -> dict[str, str]:
    """Числовой формат каждого столбца — если он один на все строки данных.

    ``NumberFormat`` диапазона отдаёт строку, когда формат у всех ячеек
    одинаковый, и пусто, когда разные. Столбец ПП обычно отформатирован
    целиком, так что одного вызова на столбец хватает; у «пёстрых» столбцов
    формат конкретной ячейки дочитывается уже при записи.
    """
    api = sheet.api
    formats: dict[str, str] = {}
    for code, column in columns.items():
        common = api.Range(
            api.Cells(first_row, column), api.Cells(last_row, column)
        ).NumberFormat
        formats[code] = common if isinstance(common, str) else ""
    return formats


def read_cells(sheet: Any, cells: list[tuple[int, int]]) -> dict[tuple[int, int], Any]:
    """Значения ячеек ``(строка, столбец)`` — одним чтением охватывающего блока.

    Нужно для проверки после записи: тысяча отдельных обращений к Excel
    стоит секунды, один прямоугольник — доли секунды.
    """
    if not cells:
        return {}
    rows = [row for row, _column in cells]
    columns = [column for _row, column in cells]
    top, left = min(rows), min(columns)
    height, width = max(rows) - top + 1, max(columns) - left + 1
    block = _as_rows(
        sheet.range((top, left), (max(rows), max(columns))).value, height, width
    )
    return {
        (row, column): from_excel(block[row - top][column - left])
        for row, column in cells
    }


def _read_codes(
    sheet: Any, config: Config, last_column: int, title: str
) -> dict[str, int]:
    """Строка кодов: ``код -> номер столбца``.

    Читается через ``Value2``: ячейка с кодом 21 может стоять в формате даты
    (в ПП 2.0 так и есть), и тогда ``Value`` вернул бы 21.01.1900.
    """
    raw = sheet.range((config.code_row, 1), (config.code_row, last_column)).api.Value2
    cells = list(raw[0]) if isinstance(raw, (list, tuple)) else [raw]

    columns: dict[str, int] = {}
    repeats: dict[str, list[str]] = defaultdict(list)
    for position, value in enumerate(cells, start=1):
        code = code_text(value)
        if not code:
            continue
        if code in columns:
            if not repeats[code]:
                repeats[code].append(get_column_letter(columns[code]))
            repeats[code].append(get_column_letter(position))
            continue
        columns[code] = position

    if not columns:
        raise WorkbookStructureError(
            f"В листе «{config.sheet_name}» ({title}) нет кодов столбцов",
            [
                (
                    f"Строка {config.code_row} пустая. Программа находит столбцы "
                    "по кодам из этой строки — проверьте настройку «строка кодов "
                    "столбцов»."
                )
            ],
        )
    if repeats:
        raise WorkbookStructureError(
            f"В листе «{config.sheet_name}» ({title}) повторяются коды столбцов",
            [
                f"Код {code} стоит в столбцах {', '.join(letters)}."
                for code, letters in list(repeats.items())[:10]
            ]
            + ["Код должен быть у каждого столбца свой — иначе столбец не найти."],
        )
    return columns


def _read_headers(
    sheet: Any, config: Config, columns: dict[str, int], width: int
) -> dict[str, str]:
    """Склеить двухуровневую шапку: ``код -> "Родитель / Уточнение"``.

    Строка над шапкой читается заодно: заголовок, объединённый по вертикали
    на несколько строк, лежит в верхней из них. Без этого такой столбец
    получал бы имя соседа слева.
    """
    top_row = config.header_row_top
    bottom_row = config.header_row_bottom
    above_row = max(top_row - 1, 1)
    block = _as_rows(
        sheet.range((above_row, 1), (bottom_row, width)).value,
        bottom_row - above_row + 1,
        width,
    )

    def text(row: int, position: int) -> str:
        return clean_str(str(_at(block[row - above_row], position) or ""))

    names: dict[int, str] = {}
    parent = ""
    for position in range(1, width + 1):
        top_text = text(top_row, position)
        bottom_text = text(bottom_row, position)
        if top_text:
            # В объединённых ячейках шапки значение лежит только в первой колонке
            parent = top_text
        elif not bottom_text and above_row < top_row and text(above_row, position):
            names[position] = text(above_row, position)
            continue
        parts = [part for part in (parent, bottom_text) if part]
        names[position] = " / ".join(dict.fromkeys(parts))
    return {code: names.get(column, "") for code, column in columns.items()}


def read_formats(
    sheet: Any, cells: list[tuple[int, int]]
) -> dict[tuple[int, int], str]:
    """Числовые форматы ячеек ``(строка, столбец)``.

    Формат нужен, чтобы заметить ячейки «Текстовый»: дату туда Excel положит
    строкой по локали США. Столбец обычно отформатирован одинаково, поэтому
    спрашиваем формат полосой — ``NumberFormat`` диапазона отдаёт строку,
    если формат один, и пусто, если разные; только тогда идём по ячейкам.
    """
    by_column: dict[int, list[int]] = defaultdict(list)
    for row, column in cells:
        by_column[column].append(row)

    result: dict[tuple[int, int], str] = {}
    for column, rows in by_column.items():
        common = sheet.range((min(rows), column), (max(rows), column)).number_format
        for row in rows:
            if isinstance(common, str):
                result[(row, column)] = common
            else:
                result[(row, column)] = str(
                    sheet.range((row, column)).number_format or ""
                )
    return result


def is_text_format(number_format: str) -> bool:
    """Ячейка отформатирована как «Текстовый»."""
    return "@" in (number_format or "")


def _at(row_values: list[Any], column: int) -> Any:
    position = column - 1
    if position < 0 or position >= len(row_values):
        return None
    return row_values[position]


def _as_rows(value: Any, rows: int, width: int) -> list[list[Any]]:
    """Прямоугольный кусок листа -> список строк.

    ``Range.formula`` отдаёт кортеж кортежей, ``Range.value`` — список списков,
    а полоса в одну строку или в один столбец приезжает плоской. Приводим всё
    к одному виду, чтобы вызывающий код не гадал: на этом программа однажды
    молча потеряла защиту формул.
    """
    if not isinstance(value, (list, tuple)):
        return [[value] * width for _ in range(rows)]
    items = list(value)
    if items and isinstance(items[0], (list, tuple)):
        return [list(item) for item in items]
    if rows == 1:
        return [items]
    if width == 1:  # вертикальная полоса
        return [[item] for item in items]
    return [items]
